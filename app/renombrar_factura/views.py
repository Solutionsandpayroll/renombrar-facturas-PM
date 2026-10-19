import io
import os
import re
import shutil
import tempfile
import zipfile
import openpyxl
from openpyxl import Workbook
from django.http import HttpResponse
from django.shortcuts import render
from .graph import buscar_pdf_por_asunto as _buscar_pdf_por_asunto
from .graph import buscar_pdf_por_nit_y_factura as _buscar_pdf_por_nit_y_factura


_DIR_TEMPORAL = os.path.join(tempfile.gettempdir(), "renombrar_facturas")


def _asegurar_dir_temporal():
    os.makedirs(_DIR_TEMPORAL, exist_ok=True)
    return _DIR_TEMPORAL


def _limpiar_temporal():
    if os.path.isdir(_DIR_TEMPORAL):
        shutil.rmtree(_DIR_TEMPORAL, ignore_errors=True)


def _ruta_pdf_temporal(nombre, indice):
    """Genera una ruta única en el directorio temporal para un PDF."""
    _asegurar_dir_temporal()
    safe = re.sub(r"[^0-9a-zA-Z._-]", "_", nombre)[:80]
    return os.path.join(_DIR_TEMPORAL, f"{indice}_{safe}")


def _solo_digitos(valor):
    if valor is None:
        return ""
    return re.sub(r"\D", "", str(valor))


def _deduplicar_por_factura(filas):
    filas_por_factura = {}
    orden_facturas = []

    for f in filas:
        factura_key = _solo_digitos(f.get("factura")) or f"__sin_factura__{f.get('documento')}"

        if factura_key not in filas_por_factura:
            filas_por_factura[factura_key] = f
            orden_facturas.append(factura_key)
        else:
            actual = filas_por_factura[factura_key]
            if not actual.get("pdf_encontrado") and f.get("pdf_encontrado"):
                filas_por_factura[factura_key] = f

    return [filas_por_factura[k] for k in orden_facturas]


def _contexto_sesion(request):
    filas = request.session.get("filas_excel", [])
    filas = _deduplicar_por_factura(filas)
    return {
        "filas": filas,
        "pdf_listo": any(f.get("pdf_encontrado") for f in filas),
    }


def inicio(request):
    return render(request, "buscar_adjuntos.html", {"error": None, **_contexto_sesion(request)})


def cargar_excel(request):
    error = None

    if request.method == "POST" and request.FILES.get("archivo_excel"):
        archivo = request.FILES["archivo_excel"]

        try:
            wb = openpyxl.load_workbook(archivo, data_only=True)
            ws = wb.active

            filas_raw = []
            fila = 6

            while True:
                documento = ws.cell(row=fila, column=1).value
                factura = ws.cell(row=fila, column=2).value
                nit_prov = ws.cell(row=fila, column=27).value
                nombre_prov = ws.cell(row=fila, column=28).value

                if documento in (None, ""):
                    break

                filas_raw.append({
                    "documento": str(documento).strip(),
                    "factura": str(factura).strip() if factura is not None else "",
                    "nit_proveedor": str(nit_prov).strip() if nit_prov is not None else "",
                    "nombre_proveedor": str(nombre_prov).strip() if nombre_prov is not None else "",
                })
                fila += 1

            if not filas_raw:
                error = "No se encontraron datos a partir de la fila 6."
                return render(request, "buscar_adjuntos.html", {"error": error, **_contexto_sesion(request)})

            _limpiar_temporal()

            unicos = {}
            for f in filas_raw:
                key = f"{_solo_digitos(f['nit_proveedor'])}_{_solo_digitos(f['factura'])}"
                if key not in unicos:
                    unicos[key] = f

            cache_por_factura = {}

            encontrados = {}
            for key, f in unicos.items():
                nit_digitos = _solo_digitos(f["nit_proveedor"])
                factura_digitos = _solo_digitos(f["factura"])
                factura_original = f["factura"].strip()

                asunto_usado = None
                pdf_path, pdf_nombre, err_busqueda = (None, None, None)

                pdf_bytes = None
                intentos = []
                if factura_digitos:
                    intentos.append(f"FEVD{factura_digitos}")
                    if factura_original != factura_digitos:
                        intentos.append(factura_original)
                    intentos.append(factura_digitos)

                if not intentos:
                    err_busqueda = "NIT y Factura vacíos, no se pudo armar ningún criterio de búsqueda."
                else:
                    errores_intentos = []
                    for criterio in intentos:
                        pdf_bytes, pdf_nombre, err_intento = _buscar_pdf_por_asunto(criterio)
                        asunto_usado = criterio
                        if pdf_bytes is not None:
                            err_busqueda = None
                            break
                        errores_intentos.append(f"'{criterio}': {err_intento}")
                    else:
                        facturas_a_probar = [x for x in (factura_original, factura_digitos) if x]
                        if nit_digitos and facturas_a_probar:
                            pdf_bytes, pdf_nombre, err_intento = _buscar_pdf_por_nit_y_factura(
                                nit_digitos, facturas_a_probar
                            )
                            if pdf_bytes is not None:
                                asunto_usado = f"NIT:{nit_digitos}"
                                err_busqueda = None
                            else:
                                err_busqueda = ("Sin coincidencias. Intentos: " + " | ".join(errores_intentos)
                                                + f" | Búsqueda por NIT: {err_intento}")
                        else:
                            err_busqueda = "Sin coincidencias. Intentos: " + " | ".join(errores_intentos)

                if pdf_bytes is not None:
                    ruta_pdf = _ruta_pdf_temporal(pdf_nombre, len(encontrados))
                    with open(ruta_pdf, "wb") as tmp:
                        tmp.write(pdf_bytes)
                    pdf_path = ruta_pdf
                    pdf_bytes = None

                encontrados[key] = {
                    "asunto_usado": asunto_usado,
                    "pdf_path": pdf_path,
                    "pdf_nombre": pdf_nombre,
                    "err_busqueda": err_busqueda,
                }

            filas_final = []
            for f in filas_raw:
                key = f"{_solo_digitos(f['nit_proveedor'])}_{_solo_digitos(f['factura'])}"
                res = encontrados[key]

                fila_resultado = dict(f)
                fila_resultado["asunto_busqueda"] = res["asunto_usado"] or ""
                if res["pdf_path"] is not None:
                    fila_resultado["pdf_encontrado"] = True
                    fila_resultado["pdf_path"] = res["pdf_path"]
                    fila_resultado["nombre_adjunto"] = res["pdf_nombre"]
                    fila_resultado["error_busqueda"] = None
                else:
                    fila_resultado["pdf_encontrado"] = False
                    fila_resultado["pdf_path"] = None
                    fila_resultado["nombre_adjunto"] = None
                    fila_resultado["error_busqueda"] = res["err_busqueda"]

                filas_final.append(fila_resultado)

            request.session["filas_excel"] = filas_final

            if not any(f["pdf_encontrado"] for f in filas_final):
                error = "Se cargó el Excel, pero no se encontró el correo/PDF para ninguna fila. Revisa el detalle por fila en la tabla."

        except Exception as e:
            error = f"Error al leer el Excel: {e}"

    return render(request, "buscar_adjuntos.html", {"error": error, **_contexto_sesion(request)})


def generar_reporte_final(request):
    error = None

    if request.method == "POST":
        filas = request.session.get("filas_excel", [])

        if not filas:
            error = "Primero debes cargar el Excel."
        elif not any(f.get("pdf_encontrado") for f in filas):
            error = "Ninguna fila tiene un PDF encontrado todavía."

        if error:
            return render(request, "buscar_adjuntos.html", {"error": error, **_contexto_sesion(request)})

        filas = _deduplicar_por_factura(filas)

        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte"
        ws.append(["Documento", "Factura", "Nit Proveedor", "Nombre Proveedor", "Factura renombrada"])
        for f in filas:
            ws.append([
                f["documento"], f["factura"], f["nit_proveedor"], f["nombre_proveedor"],
                "Sí" if f.get("pdf_encontrado") else "No",
            ])

        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        zip_buffer = io.BytesIO()
        nombres_usados = {}
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
            for f in filas:
                if not f.get("pdf_encontrado"):
                    continue

                with open(f["pdf_path"], "rb") as src:
                    pdf_bytes = src.read()

                nombre_pdf = f"{f['documento']}-{f['nombre_proveedor']}-{f['factura']}.pdf"
                nombre_pdf = nombre_pdf.replace("/", "-").replace("\\", "-")

                if nombre_pdf in nombres_usados:
                    nombres_usados[nombre_pdf] += 1
                    base, ext = nombre_pdf.rsplit(".", 1)
                    nombre_pdf = f"{base} ({nombres_usados[nombre_pdf]}).{ext}"
                else:
                    nombres_usados[nombre_pdf] = 0

                zf.writestr(nombre_pdf, pdf_bytes)

            zf.writestr("reporte.xlsx", excel_buffer.getvalue())

        zip_buffer.seek(0)

        request.session.pop("filas_excel", None)
        _limpiar_temporal()

        response_http = HttpResponse(zip_buffer.getvalue(), content_type="application/zip")
        response_http["Content-Disposition"] = 'attachment; filename="resultado.zip"'
        return response_http

    return render(request, "buscar_adjuntos.html", {"error": error, **_contexto_sesion(request)})